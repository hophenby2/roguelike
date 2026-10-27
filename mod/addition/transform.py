from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Font
from collections import OrderedDict
from lupa import LuaRuntime
import re
import json

lua = LuaRuntime(unpack_returned_tuples=True)


def convert_lua_to_dict_from_file(lua_code):
    # 正则表达式动态匹配 Lua 代码，捕获所有可能的字段（包括可选字段）
    pattern = r'\[(\d+)\] = {\s*(.*?)\s*},'

    matches = re.findall(pattern, lua_code, re.DOTALL)

    lua_data = {}

    for match in matches:
        # match[0] 是 ID，match[1] 是字段内容
        id_ = int(match[0])
        fields_str = match[1]

        # 初始化默认字段
        entry = {
            "name": None,
            "stitle": None,
            "sdes": None,
            "detail": None,
            "repeatd": None,
            "collabd": None,
            "condd": None,
            "unlockd": None,
            "initialTake": None,
            "tags": None,
            "other": None
        }

        # 匹配每个字段，更新字典
        field_patterns = [
            ("name", r'name\s*=\s*\[\[([^]]+)\]\]', 0),
            ("stitle", r'stitle\s*=\s*\[\[([^]]+)\]\]', 0),
            ("sdes", r'sdes\s*=\s*\[\[([^]]+)\]\]', 0),
            ("detail", r'detail\s*=\s*\[\[([^]]+)\]\]', 0),
            ("repeatd", r'repeatd\s*=\s*\[\[([^]]+)\]\]', 0),
            ("collabd", r'collabd\s*=\s*\[\[([^]]+)\]\]', 0),
            ("unlockd", r'unlockd\s*=\s*\[\[([^]]+)\]\]', 0),
            ("condd", r'condd\s*=\s*\[\[([^]]+)\]\]', 0),
            ("tags", r'tags\s*=\s*(.+)', 0),
            ("initialTake", r'initialTake\s*=\s*(true|false)', 0),
            ("other", r'other\s*=\s*\[\[([^]]+)\]\]', 0)
        ]

        for field_name, pattern, mode in field_patterns:
            match_result = re.search(pattern, fields_str, mode)
            if match_result:
                value = match_result.group(1)
                if field_name == "tags":
                    if value == "nil":
                        entry["tags"] = []
                    elif "{" in value:
                        # 处理多个标签
                        entry["tags"] = value + " }"
                    elif value:
                        # 处理单个标签
                        entry["tags"] = value[:-1]
                elif field_name == "initialTake":
                    entry[field_name] = value == "true"
                else:
                    entry[field_name] = value

        # 将该条目加入字典
        lua_data[id_] = entry

    return lua_data


row_index = (
    'id', 'name', 'stitle', 'sdes', 'detail', 'repeatd', 'collabd', 'unlockd', 'condd', 'tags', 'initialTake', 'other')
row_map = {
    'id': 0,
    'name': 1,
    'stitle': 2,
    'sdes': 3,
    'detail': 4,
    'repeatd': 5,
    'collabd': 6,
    'unlockd': 7,
    'condd': 8,
    'tags': 9,
    'initialTake': 10,
    'other': 11,
}
row_dict = {
    'id': 'ID',
    'name': '名称',
    'stitle': '副标题',
    'sdes': '简述',
    'detail': '详细描述',
    'repeatd': '叠加效果描述',
    'collabd': '道具协同描述',
    'unlockd': '解锁方式描述',
    'condd': '出现前提描述',
    'tags': '标签',
    'initialTake': '初始携带',
    'other': '其他',
}
width_data = [6, 30, 30, 100, 100, 100, 100, 100, 100, 30, 20, 30]
align_data = ['center', 'center', 'left', 'left', 'left', 'left', 'left', 'left', 'left', 'left', 'left', 'left', 'left']

f = open('addition_des_en.lua', 'r', encoding='utf-8')
data_eng = lua.execute(f.read())
f.close()
letter = ('A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q', 'R', 'S', 'T', 'U', 'V', 'W', 'X', 'Y', 'Z')
wb = Workbook()
wb.remove(wb.active)
sheet_name = ('0', '1', '2', '3', '4')
for i in range(5):
    ws = wb.create_sheet(sheet_name[i])

    for k in range(12):
        ws[f'{letter[k]}{1}'] = row_dict[row_index[k]]
        ws[f'{letter[k]}{1}'].alignment = Alignment(horizontal='center', vertical='center')
        ws[f'{letter[k]}{1}'].font = Font(bold=True, size=14, name='黑体')
        ws.column_dimensions[f'{letter[k]}'].width = 8

    with open(f'{i}.lua', 'r', encoding='utf-8') as f:
        text = f.read()
        result = convert_lua_to_dict_from_file(text)
        ptr = 2

        for k, v in result.items():
            ws[f'A{ptr}'] = k
            ws[f'A{ptr}'].font = Font(bold=True, size=15, name='黑体')
            ws[f'A{ptr}'].alignment = Alignment(horizontal='center', vertical='center')
            ws.merge_cells(f'A{ptr}:A{ptr + 1}')
            for kk, vv in v.items():

                pos1 = f'{letter[row_map[kk]]}{ptr}'
                pos2 = f'{letter[row_map[kk]]}{ptr + 1}'

                vv2 = ""
                if data_eng[i + 1] and data_eng[i + 1][k] and data_eng[i + 1][k][kk]:
                    if type(data_eng[i + 1][k][kk]) == str:
                        vv2 = data_eng[i + 1][k][kk]
                        ws[pos2] = data_eng[i + 1][k][kk]
                    else:
                        str_list = []
                        for key, value in data_eng[i + 1][k][kk].items():
                            str_list.append(value)
                        vv2 = ', '.join(str_list)
                        ws[pos2] = ', '.join(str_list)
                ws[pos1] = vv
                ws[pos2] = vv2
                ws.column_dimensions[f'{letter[row_map[kk]]}'].width = width_data[row_map[kk]]

                if ws.row_dimensions[ptr].height:
                    ws.row_dimensions[ptr].height = max(ws.row_dimensions[ptr].height, isinstance(vv, str) and (vv.count('\n') * 20+20) or 20)
                else:
                    ws.row_dimensions[ptr].height = 15
                if ws.row_dimensions[ptr + 1].height:
                    ws.row_dimensions[ptr + 1].height = max(ws.row_dimensions[ptr + 1].height, isinstance(vv2, str) and (vv2.count('\n') * 20+20) or 20)
                else:
                    ws.row_dimensions[ptr + 1].height = 15
                ws[pos1].alignment = Alignment(wrapText=True, horizontal=align_data[row_map[kk]], vertical='center')
                ws[pos2].alignment = Alignment(wrapText=True, horizontal=align_data[row_map[kk]], vertical='center')
                ws[pos1].font = Font(size=14, name='黑体')
                ws[pos2].font = Font(size=14, name='Arial')

            ptr += 2
    ##print(json.dumps(result, indent=2, ensure_ascii=False))
    f.close()

wb.save("example.xlsx")
